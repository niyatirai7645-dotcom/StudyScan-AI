import time
import re
from io import BytesIO
from datetime import datetime

import streamlit as st
from PIL import Image

from ai import generate_response, ask_studyscan
from prompts import get_prompt
from pdf_utils import (
    get_pdf_page_count,
    pdf_to_image_batches
)



from auth import (
    sign_up,
    login,
    logout,
    get_current_user,
    forgot_password,
    update_password,
    verify_password_recovery
)

from chat_db import (
    load_user_chats,
    save_chat_to_database,
    delete_chat_from_database,
    delete_all_user_chats
)


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="StudyScan AI",
    page_icon="📚",
    layout="wide"
)


# ============================================================
# CUSTOM CSS
# ============================================================

st.markdown(
    """
    <style>

    /* =========================================================
       STUDYSCAN GLOBAL DESIGN SYSTEM
       ========================================================= */

    :root {
        --ss-bg: #07111F;
        --ss-bg-2: #0A1628;
        --ss-sidebar: #081525;
        --ss-card: #0D1B2E;
        --ss-card-2: #10233B;
        --ss-border: #1D3553;
        --ss-border-light: #274565;

        --ss-text: #F7FAFC;
        --ss-text-soft: #B8C7DA;
        --ss-text-muted: #8193AA;

        --ss-primary: #3B82F6;
        --ss-primary-hover: #2563EB;
        --ss-primary-soft: rgba(59, 130, 246, 0.14);

        --ss-success: #22C55E;
        --ss-radius: 14px;
        --ss-radius-small: 10px;
    }


    /* =========================================================
       GLOBAL PAGE
       ========================================================= */

    html,
    body,
    [data-testid="stAppViewContainer"] {

        background:
            radial-gradient(
                circle at 75% 5%,
                rgba(37, 99, 235, 0.08),
                transparent 30%
            ),
            var(--ss-bg) !important;

        color: var(--ss-text) !important;
    }


    [data-testid="stHeader"] {

        background: transparent !important;
    }


    .main {

        background: transparent !important;
    }


    .main .block-container {

        max-width: 1500px !important;

        padding-top: 2.5rem !important;
        padding-bottom: 100px !important;
        padding-left: 3rem !important;
        padding-right: 3rem !important;
    }


    /* =========================================================
       REMOVE STREAMLIT DEFAULT DECORATION
       ========================================================= */

    [data-testid="stDecoration"] {

        display: none !important;
    }


    /* =========================================================
       TYPOGRAPHY
       ========================================================= */

    .main h1 {

        color: var(--ss-text) !important;

        font-size: 2.45rem !important;
        font-weight: 750 !important;

        letter-spacing: -0.035em !important;

        margin-bottom: 0.35rem !important;
    }


    .main h2 {

        color: var(--ss-text) !important;

        font-size: 1.55rem !important;
        font-weight: 700 !important;

        letter-spacing: -0.02em !important;
    }


    .main h3 {

        color: var(--ss-text) !important;

        font-size: 1.2rem !important;
        font-weight: 650 !important;
    }


    .main p {

        color: var(--ss-text-soft);
    }


    /* =========================================================
       SIDEBAR
       ========================================================= */

    section[data-testid="stSidebar"] {

        background:
            linear-gradient(
                180deg,
                #081525 0%,
                #07111F 100%
            ) !important;

        border-right: 1px solid var(--ss-border) !important;
    }


    section[data-testid="stSidebar"] > div {

        background: transparent !important;

        padding-top: 1.25rem !important;
    }


    section[data-testid="stSidebar"] h1,
    section[data-testid="stSidebar"] h2,
    section[data-testid="stSidebar"] h3 {

        color: var(--ss-text) !important;
    }


    section[data-testid="stSidebar"] p {

        color: var(--ss-text-soft) !important;
    }


    /* Sidebar buttons */

    section[data-testid="stSidebar"] button {

        border-radius: 11px !important;

        border: 1px solid transparent !important;

        background: transparent !important;

        color: var(--ss-text-soft) !important;

        transition:
            background 0.18s ease,
            border-color 0.18s ease,
            transform 0.18s ease !important;
    }


    section[data-testid="stSidebar"] button:hover {

        background: rgba(59, 130, 246, 0.10) !important;

        border-color: rgba(59, 130, 246, 0.18) !important;

        color: var(--ss-text) !important;

        transform: translateX(2px);
    }


    /* =========================================================
       TOP RIGHT MENU
       ========================================================= */

    .st-key-studyscan_top_menu {

        position: fixed !important;

        top: 16px !important;
        right: 22px !important;

        z-index: 1000000 !important;
    }


    .st-key-studyscan_top_menu > button {

        width: 44px !important;
        min-width: 44px !important;

        height: 44px !important;
        min-height: 44px !important;

        padding: 0 !important;

        border-radius: 50% !important;

        background: rgba(13, 27, 46, 0.92) !important;

        border: 1px solid var(--ss-border-light) !important;

        color: var(--ss-text) !important;

        font-size: 23px !important;

        box-shadow:
            0 8px 25px rgba(0, 0, 0, 0.28) !important;

        transition:
            background 0.18s ease,
            border-color 0.18s ease,
            transform 0.18s ease !important;
    }


    .st-key-studyscan_top_menu > button:hover {

        background: var(--ss-card-2) !important;

        border-color: #3B82F6 !important;

        transform: scale(1.05);
    }


    /* Popover itself */

    div[data-testid="stPopoverBody"] {

        background: #0D1B2E !important;

        border: 1px solid var(--ss-border-light) !important;

        border-radius: 16px !important;

        box-shadow:
            0 24px 60px rgba(0, 0, 0, 0.45) !important;
    }


    /* =========================================================
       MAIN HERO / HEADER
       ========================================================= */

    .main h1:first-of-type {

        margin-top: 0.5rem !important;
    }


    .main h1:first-of-type + div p {

        color: var(--ss-text-soft) !important;

        font-size: 1.02rem !important;

        margin-top: -0.15rem !important;

        margin-bottom: 1.4rem !important;
    }


    /* =========================================================
       DIVIDERS
       ========================================================= */

    .main hr {

        border: none !important;

        border-top: 1px solid var(--ss-border) !important;

        margin: 1.35rem 0 !important;

        opacity: 0.85 !important;
    }


    /* =========================================================
       STANDARD BUTTONS
       ========================================================= */

    .main button {

        border-radius: 10px !important;

        font-weight: 600 !important;

        transition:
            background 0.18s ease,
            border-color 0.18s ease,
            transform 0.18s ease,
            box-shadow 0.18s ease !important;
    }


    .main button:hover {

        transform: translateY(-1px);
    }


    /* Primary buttons */

    .main button[kind="primary"] {

        background:
            linear-gradient(
                135deg,
                #3B82F6,
                #2563EB
            ) !important;

        color: #FFFFFF !important;

        border: 1px solid #3B82F6 !important;

        box-shadow:
            0 8px 25px rgba(37, 99, 235, 0.20) !important;
    }


    .main button[kind="primary"]:hover {

        background:
            linear-gradient(
                135deg,
                #60A5FA,
                #3B82F6
            ) !important;

        box-shadow:
            0 12px 30px rgba(37, 99, 235, 0.30) !important;
    }


    /* Secondary buttons */

    .main button[kind="secondary"] {

        background: var(--ss-card) !important;

        color: var(--ss-text) !important;

        border: 1px solid var(--ss-border-light) !important;
    }


    .main button[kind="secondary"]:hover {

        background: var(--ss-card-2) !important;

        border-color: #3B82F6 !important;
    }


    /* =========================================================
       INPUTS
       ========================================================= */

    .main input,
    .main textarea {

        background: #0B192B !important;

        color: var(--ss-text) !important;

        border: 1px solid var(--ss-border) !important;

        border-radius: var(--ss-radius-small) !important;
    }


    .main input:focus,
    .main textarea:focus {

        border-color: #3B82F6 !important;

        box-shadow:
            0 0 0 1px #3B82F6,
            0 0 18px rgba(59, 130, 246, 0.12) !important;
    }


    .main input::placeholder,
    .main textarea::placeholder {

        color: #687D96 !important;
    }


    /* =========================================================
       SELECTBOXES
       ========================================================= */

    .main [data-baseweb="select"] > div {

        background: #0B192B !important;

        border-color: var(--ss-border) !important;

        border-radius: var(--ss-radius-small) !important;

        color: var(--ss-text) !important;
    }


    .main [data-baseweb="select"] > div:hover {

        border-color: #3B82F6 !important;
    }


    [data-baseweb="popover"] {

        background: #0D1B2E !important;

        border: 1px solid var(--ss-border-light) !important;
    }


    [role="option"] {

        background: #0D1B2E !important;

        color: var(--ss-text) !important;
    }


    [role="option"]:hover {

        background: rgba(59, 130, 246, 0.12) !important;
    }


    /* =========================================================
       RADIO / MODE SELECTOR
       ========================================================= */

    .main [role="radiogroup"] {

        gap: 10px !important;
    }


    .main [role="radiogroup"] label {

        background: var(--ss-card) !important;

        border: 1px solid var(--ss-border) !important;

        border-radius: 12px !important;

        padding: 11px 18px !important;

        color: var(--ss-text-soft) !important;

        transition:
            background 0.18s ease,
            border-color 0.18s ease,
            color 0.18s ease !important;
    }


    .main [role="radiogroup"] label:hover {

        background: var(--ss-card-2) !important;

        border-color: #35577D !important;

        color: var(--ss-text) !important;
    }


    .main [role="radiogroup"] label:has(input:checked) {

        background:
            rgba(59, 130, 246, 0.14) !important;

        border-color:
            rgba(59, 130, 246, 0.75) !important;

        color: #FFFFFF !important;
    }


    /* =========================================================
       FILE UPLOADER
       ========================================================= */

    [data-testid="stFileUploader"] {

        margin-top: 0.5rem !important;
    }


    [data-testid="stFileUploaderDropzone"] {

        background:
            linear-gradient(
                145deg,
                rgba(13, 27, 46, 0.96),
                rgba(10, 22, 40, 0.96)
            ) !important;

        border: 1px dashed #35577D !important;

        border-radius: 16px !important;

        padding: 28px 20px !important;

        transition:
            border-color 0.2s ease,
            background 0.2s ease,
            transform 0.2s ease !important;
    }


    [data-testid="stFileUploaderDropzone"]:hover {

        border-color: #3B82F6 !important;

        background:
            linear-gradient(
                145deg,
                rgba(17, 39, 66, 0.98),
                rgba(10, 27, 48, 0.98)
            ) !important;

        transform: translateY(-1px);
    }


    [data-testid="stFileUploaderDropzoneInstructions"] {

        color: var(--ss-text-soft) !important;
    }


    [data-testid="stFileUploaderDropzoneInstructions"] small {

        color: var(--ss-text-muted) !important;
    }


    /* =========================================================
       CARDS / VERTICAL BLOCKS
       ========================================================= */

    .main [data-testid="stVerticalBlockBorderWrapper"] {

        background: rgba(13, 27, 46, 0.65) !important;

        border: 1px solid var(--ss-border) !important;

        border-radius: 16px !important;
    }


    /* =========================================================
       SUCCESS / WARNING / ERROR / INFO
       ========================================================= */

    .main [data-testid="stAlert"] {

        border-radius: 12px !important;

        border: 1px solid var(--ss-border) !important;

        background: var(--ss-card) !important;
    }


    /* =========================================================
       SPINNER
       ========================================================= */

    .main [data-testid="stSpinner"] {

        color: #60A5FA !important;
    }


    /* =========================================================
       DOWNLOAD BUTTONS
       ========================================================= */

    .main [data-testid="stDownloadButton"] button {

        background: #0D1B2E !important;

        color: var(--ss-text) !important;

        border: 1px solid var(--ss-border-light) !important;

        border-radius: 10px !important;
    }


    .main [data-testid="stDownloadButton"] button:hover {

        background: #122742 !important;

        border-color: #3B82F6 !important;
    }


    /* =========================================================
       ASK STUDYSCAN — FIXED RIGHT PANEL
       ========================================================= */

    .st-key-ask_panel {

        position: fixed !important;

        top: 0 !important;
        right: 0 !important;

        width: 410px !important;
        max-width: 410px !important;

        height: 100vh !important;

        z-index: 999999 !important;

        background:
            linear-gradient(
                180deg,
                #0B1F3A 0%,
                #08192D 100%
            ) !important;

        border-left: 1px solid #254667 !important;

        box-shadow:
            -18px 0 50px rgba(0, 0, 0, 0.35) !important;

        padding: 24px !important;

        box-sizing: border-box !important;

        overflow-y: auto !important;

        overflow-x: hidden !important;
    }


    /* Keep the internal panel background consistent */

    .st-key-ask_panel,
    .st-key-ask_panel > div,
    .st-key-ask_panel > div > div,
    .st-key-ask_panel [data-testid="stVerticalBlock"],
    .st-key-ask_panel [data-testid="stElementContainer"] {

        background-color: transparent !important;
    }


    /* Main page gives the panel its own space */

    body:has(.st-key-ask_panel)
    .main
    .block-container {

        padding-right: 435px !important;
    }


    /* Ask header */

    .ask-header-title {

        font-size: 25px !important;

        font-weight: 750 !important;

        color: #FFFFFF !important;

        letter-spacing: -0.02em !important;

        margin-bottom: 3px !important;
    }


    .ask-header-subtitle {

        font-size: 14px !important;

        color: #AFC1D6 !important;

        margin-bottom: 12px !important;
    }


    /* Ask chat area */

    .ask-chat-area {

        width: 100% !important;

        overflow-x: hidden !important;

        word-wrap: break-word !important;

        overflow-wrap: anywhere !important;

        padding-bottom: 16px !important;
    }


    /* User message */

    .ask-user-row {

        width: 100% !important;

        display: flex !important;

        justify-content: flex-end !important;

        margin: 10px 0 !important;
    }


    .ask-user-bubble {

        max-width: 82% !important;

        background:
            linear-gradient(
                135deg,
                #3B82F6,
                #2563EB
            ) !important;

        color: #FFFFFF !important;

        border-radius:
            18px 18px 5px 18px !important;

        padding: 12px 15px !important;

        font-size: 14px !important;

        line-height: 1.55 !important;

        box-shadow:
            0 8px 22px rgba(0, 0, 0, 0.18) !important;
    }


    /* Bot message */

    .ask-bot-row {

        width: 100% !important;

        display: flex !important;

        justify-content: flex-start !important;

        margin: 10px 0 !important;
    }


    .ask-bot-bubble {

        max-width: 84% !important;

        background: #132B49 !important;

        color: #F5F7FA !important;

        border: 1px solid #294B70 !important;

        border-radius:
            18px 18px 18px 5px !important;

        padding: 12px 15px !important;

        font-size: 14px !important;

        line-height: 1.55 !important;
    }


    .ask-user-label,
    .ask-bot-label {

        font-size: 10px !important;

        font-weight: 750 !important;

        letter-spacing: 0.04em !important;

        text-transform: uppercase !important;

        margin-bottom: 5px !important;
    }


    .ask-user-label {

        color: #DCEBFF !important;
    }


    .ask-bot-label {

        color: #8DB9E8 !important;
    }


    /* Ask textbox */

    .st-key-ask_question textarea {

        min-height: 140px !important;

        height: 140px !important;

        resize: vertical !important;

        background: #08192D !important;

        color: #FFFFFF !important;

        caret-color: #FFFFFF !important;

        border: 1px solid #315578 !important;

        border-radius: 12px !important;
    }


    .st-key-ask_question textarea:focus {

        border-color: #3B82F6 !important;

        box-shadow:
            0 0 0 1px #3B82F6,
            0 0 20px rgba(59, 130, 246, 0.12) !important;
    }


    .st-key-ask_question textarea::placeholder {

        color: #6F86A1 !important;
    }


    /* Ask send */

    .st-key-ask_question_form button {

        background:
            linear-gradient(
                135deg,
                #3B82F6,
                #2563EB
            ) !important;

        color: #FFFFFF !important;

        border: none !important;

        border-radius: 10px !important;

        min-height: 44px !important;

        font-weight: 650 !important;
    }


    .st-key-ask_question_form button:hover {

        box-shadow:
            0 8px 25px rgba(37, 99, 235, 0.28) !important;
    }


    /* Ask controls */

    .st-key-ask_clear button,
    .st-key-ask_close button {

        background: #122942 !important;

        color: #FFFFFF !important;

        border: 1px solid #315578 !important;

        border-radius: 10px !important;

        min-height: 42px !important;
    }


    .st-key-ask_clear button:hover,
    .st-key-ask_close button:hover {

        background: #183754 !important;

        border-color: #3B82F6 !important;
    }


    .st-key-ask_panel hr {

        border-color: #294B70 !important;
    }


    .st-key-ask_panel p,
    .st-key-ask_panel div,
    .st-key-ask_panel span {

        overflow-wrap: anywhere !important;

        word-break: break-word !important;
    }


    /* =========================================================
       SCROLLBAR
       ========================================================= */

    ::-webkit-scrollbar {

        width: 7px;
        height: 7px;
    }


    ::-webkit-scrollbar-track {

        background: transparent;
    }


    ::-webkit-scrollbar-thumb {

        background: #29435F;

        border-radius: 20px;
    }


    ::-webkit-scrollbar-thumb:hover {

        background: #3B82F6;
    }


    /* =========================================================
       MOBILE
       ========================================================= */

    @media (max-width: 850px) {

        .main .block-container {

            padding-left: 1rem !important;

            padding-right: 1rem !important;
        }


        .main h1 {

            font-size: 2rem !important;
        }


        .st-key-ask_panel {

            width: 100vw !important;

            max-width: 100vw !important;

            padding: 20px !important;
        }


        body:has(.st-key-ask_panel)
        .main
        .block-container {

            padding-right: 1rem !important;
        }


        .main [role="radiogroup"] {

            flex-wrap: wrap !important;
        }
    }

    /* =========================================================
   BATCH 2 — AUTHENTICATION EXPERIENCE
   ========================================================= */


/* =========================================================
   AUTH PAGE BACKGROUND
   ========================================================= */

body:has(.auth-shell) {

    background:
        radial-gradient(
            circle at 15% 20%,
            rgba(59, 130, 246, 0.12),
            transparent 30%
        ),
        radial-gradient(
            circle at 85% 80%,
            rgba(37, 99, 235, 0.10),
            transparent 30%
        ),
        #07111F !important;
}


body:has(.auth-shell)
section[data-testid="stSidebar"] {

    display: none !important;
}


body:has(.auth-shell)
[data-testid="stHeader"] {

    display: none !important;
}


body:has(.auth-shell)
.main .block-container {

    max-width: 1250px !important;

    min-height: 100vh !important;

    padding:
        30px
        40px
        50px
        40px !important;

    display: flex !important;

    align-items: center !important;

    justify-content: center !important;
}


/* =========================================================
   AUTH OUTER CONTAINER
   ========================================================= */

.auth-shell {

    width: 100%;

    max-width: 1080px;

    min-height: 650px;

    display: grid;

    grid-template-columns:
        1.05fr
        0.95fr;

    background:
        linear-gradient(
            145deg,
            rgba(13, 27, 46, 0.94),
            rgba(7, 18, 32, 0.98)
        );

    border:
        1px solid rgba(83, 125, 169, 0.25);

    border-radius: 28px;

    overflow: hidden;

    box-shadow:
        0 35px 100px rgba(0, 0, 0, 0.45);
}


/* =========================================================
   LEFT BRAND PANEL
   ========================================================= */

.auth-brand {

    position: relative;

    padding: 55px;

    display: flex;

    flex-direction: column;

    justify-content: space-between;

    overflow: hidden;

    background:
        radial-gradient(
            circle at 25% 20%,
            rgba(59, 130, 246, 0.20),
            transparent 35%
        ),
        linear-gradient(
            145deg,
            #0C2440,
            #081525
        );
}


.auth-brand::before {

    content: "";

    position: absolute;

    width: 280px;
    height: 280px;

    border-radius: 50%;

    right: -100px;
    top: -100px;

    background:
        rgba(59, 130, 246, 0.10);

    filter: blur(5px);
}


.auth-brand::after {

    content: "";

    position: absolute;

    width: 220px;
    height: 220px;

    border-radius: 50%;

    left: -110px;
    bottom: -100px;

    background:
        rgba(37, 99, 235, 0.08);

    filter: blur(4px);
}


/* =========================================================
   BRAND CONTENT
   ========================================================= */

.auth-brand-content {

    position: relative;

    z-index: 2;
}


.auth-logo {

    display: flex;

    align-items: center;

    gap: 12px;

    font-size: 24px;

    font-weight: 750;

    color: #FFFFFF;

    letter-spacing: -0.025em;
}


.auth-logo-icon {

    width: 44px;
    height: 44px;

    display: flex;

    align-items: center;
    justify-content: center;

    border-radius: 13px;

    background:
        linear-gradient(
            135deg,
            #3B82F6,
            #2563EB
        );

    box-shadow:
        0 10px 30px rgba(37, 99, 235, 0.30);

    font-size: 21px;
}


.auth-brand h1 {

    margin-top: 65px;

    margin-bottom: 18px;

    font-size: 44px !important;

    line-height: 1.08;

    letter-spacing: -0.045em;

    color: #FFFFFF !important;
}


.auth-brand h1 span {

    color: #60A5FA;
}


.auth-brand-description {

    max-width: 450px;

    font-size: 16px;

    line-height: 1.7;

    color: #AFC1D6;
}


/* =========================================================
   FEATURE PILLS
   ========================================================= */

.auth-feature-list {

    position: relative;

    z-index: 2;

    display: flex;

    flex-wrap: wrap;

    gap: 9px;

    margin-top: 35px;
}


.auth-feature {

    padding:
        8px
        12px;

    border:
        1px solid rgba(100, 150, 200, 0.20);

    border-radius: 999px;

    background:
        rgba(255, 255, 255, 0.035);

    color: #C6D5E6;

    font-size: 12px;
}


/* =========================================================
   BRAND FOOTER
   ========================================================= */

.auth-brand-footer {

    position: relative;

    z-index: 2;

    color: #6F86A1;

    font-size: 12px;
}


/* =========================================================
   LOGIN AREA
   ========================================================= */

.auth-form-area {

    padding:
        55px
        50px;

    display: flex;

    flex-direction: column;

    justify-content: center;

    background:
        rgba(6, 16, 29, 0.72);

    border-left:
        1px solid rgba(83, 125, 169, 0.16);
}


.auth-form-title {

    font-size: 29px;

    font-weight: 750;

    color: #FFFFFF;

    letter-spacing: -0.035em;

    margin-bottom: 8px;
}


.auth-form-subtitle {

    color: #8193AA;

    font-size: 14px;

    line-height: 1.5;

    margin-bottom: 28px;
}


/* =========================================================
   AUTH TABS
   ========================================================= */

.auth-tabs {

    display: grid;

    grid-template-columns: 1fr 1fr;

    gap: 5px;

    padding: 5px;

    margin-bottom: 24px;

    background: #09182A;

    border:
        1px solid #1C3552;

    border-radius: 12px;
}


.auth-tabs > div {

    text-align: center;

    padding: 9px;

    border-radius: 8px;

    color: #8193AA;

    font-size: 13px;

    font-weight: 600;
}


.auth-tabs .active {

    background: #132B49;

    color: #FFFFFF;

    box-shadow:
        0 3px 10px rgba(0, 0, 0, 0.20);
}


/* =========================================================
   AUTH INPUTS
   ========================================================= */

body:has(.auth-shell)
input {

    background:
        #09182A !important;

    color:
        #FFFFFF !important;

    border:
        1px solid #1D3A59 !important;

    border-radius:
        11px !important;

    min-height:
        46px !important;
}


body:has(.auth-shell)
input:hover {

    border-color:
        #315578 !important;
}


body:has(.auth-shell)
input:focus {

    border-color:
        #3B82F6 !important;

    box-shadow:
        0 0 0 1px #3B82F6,
        0 0 18px rgba(59, 130, 246, 0.12) !important;
}


body:has(.auth-shell)
input::placeholder {

    color:
        #637A94 !important;
}


/* =========================================================
   AUTH LABELS
   ========================================================= */

body:has(.auth-shell)
label {

    color:
        #B8C7DA !important;

    font-size:
        13px !important;

    font-weight:
        550 !important;
}


/* =========================================================
   AUTH PRIMARY BUTTON
   ========================================================= */

body:has(.auth-shell)
button[kind="primary"] {

    min-height:
        48px !important;

    border-radius:
        11px !important;

    background:
        linear-gradient(
            135deg,
            #3B82F6,
            #2563EB
        ) !important;

    border:
        1px solid #3B82F6 !important;

    color:
        #FFFFFF !important;

    font-weight:
        650 !important;

    box-shadow:
        0 10px 28px rgba(37, 99, 235, 0.22) !important;
}


body:has(.auth-shell)
button[kind="primary"]:hover {

    background:
        linear-gradient(
            135deg,
            #60A5FA,
            #3B82F6
        ) !important;

    box-shadow:
        0 14px 34px rgba(37, 99, 235, 0.32) !important;
}


/* =========================================================
   AUTH SECONDARY BUTTONS
   ========================================================= */

body:has(.auth-shell)
button[kind="secondary"] {

    color:
        #9DB0C5 !important;

    background:
        transparent !important;

    border:
        1px solid transparent !important;
}


body:has(.auth-shell)
button[kind="secondary"]:hover {

    color:
        #FFFFFF !important;

    background:
        rgba(59, 130, 246, 0.08) !important;
}


/* =========================================================
   AUTH DIVIDERS
   ========================================================= */

body:has(.auth-shell)
hr {

    border-color:
        #1C3552 !important;
}


/* =========================================================
   AUTH ALERTS
   ========================================================= */

body:has(.auth-shell)
[data-testid="stAlert"] {

    border-radius:
        11px !important;

    background:
        #0D1B2E !important;

    border:
        1px solid #294563 !important;
}


/* =========================================================
   AUTH RESPONSIVE
   ========================================================= */

@media (max-width: 850px) {

    body:has(.auth-shell)
    .main .block-container {

        padding:
            20px !important;
    }


    .auth-shell {

        grid-template-columns: 1fr;

        min-height: auto;

        max-width: 600px;
    }


    .auth-brand {

        padding: 35px;

        min-height: 360px;
    }


    .auth-brand h1 {

        margin-top: 35px;

        font-size: 34px !important;
    }


    .auth-form-area {

        padding:
            35px 28px;
    }
}


@media (max-width: 520px) {

    .auth-brand {

        padding: 28px;
    }


    .auth-form-area {

        padding:
            28px 22px;
    }


    .auth-brand h1 {

        font-size: 30px !important;
    }
}

/* Product workspace visual system: calm ink surfaces with one cobalt accent. */
:root {
    --ss-bg: #0D1726;
    --ss-card: #162235;
    --ss-card-2: #1B2A40;
    --ss-border: rgba(148, 163, 184, 0.20);
    --ss-border-light: rgba(148, 163, 184, 0.32);
    --ss-primary: #3B82F6;
    --ss-primary-hover: #60A5FA;
    --ss-text-soft: #B8C7DA;
}

[data-testid="stAppViewContainer"] {
    background: #0D1726 !important;
}

.main .block-container { max-width: 1220px !important; padding-top: 1.5rem !important; }
.main h1 { font-size: clamp(2rem, 3vw, 2.8rem) !important; line-height: 1.14 !important; }
.eyebrow { color: #78B4FF !important; font-size: .72rem !important; font-weight: 700 !important; letter-spacing: .10em !important; margin: 0 0 .55rem !important; }

.app-hero, .auth-hero {
    display: block; padding: clamp(1.4rem, 3vw, 2.4rem); margin-bottom: 1.75rem;
    background: #162235; border: 1px solid var(--ss-border); border-left: 3px solid #3B82F6;
    border-radius: 16px; box-shadow: 0 16px 36px rgba(6, 13, 24, .18);
}
.app-hero h1, .auth-hero h1 { margin: 0 0 .65rem !important; }
.app-hero p:not(.eyebrow), .auth-hero p:not(.eyebrow) { max-width: 640px; margin: 0 !important; color: #B8C7DA !important; font-size: 1rem !important; }
.hero-orb { display: none; }
.brand-mark { display:grid; place-items:center; float:left; width:42px; height:42px; margin: .15rem 1rem 2rem 0; border-radius:12px; color:#F3F7FC; font-size:1rem; font-weight:800; background:#2563EB; }
.auth-hero { max-width: 720px; margin: 5vh auto 1.25rem; }
body:has(.auth-page-marker) .main .block-container { max-width: 720px !important; }
body:has(.auth-page-marker) [data-testid="stTabs"] { padding: 1.4rem; border: 1px solid var(--ss-border); border-radius: 16px; background: #162235; box-shadow: 0 16px 36px rgba(6,13,24,.18); }

.sidebar-brand { display:flex; align-items:center; gap:.65rem; font-size:1.2rem; margin-bottom:.15rem; }
.sidebar-brand span { display:grid; place-items:center; width:32px; height:32px; border-radius:10px; background:#2563EB; color:#F3F7FC; font-size:.85rem; }
section[data-testid="stSidebar"] { background: #111D2E !important; }
section[data-testid="stSidebar"] [data-testid="stButton"] button { text-align:left !important; }

[data-testid="stFileUploaderDropzone"] { border-color: rgba(96, 165, 250, .46) !important; background: #111D2E !important; border-radius: 14px !important; }
[data-testid="stFileUploaderDropzone"]:hover { border-color: #60A5FA !important; background: #162235 !important; }
.main [data-baseweb="select"] > div, .main input, .main textarea { background: #111D2E !important; border-color: var(--ss-border) !important; }
.main button[kind="primary"] { background: #2563EB !important; border-color: #2563EB !important; box-shadow: none !important; }
.main button[kind="primary"]:hover { background: #3B82F6 !important; border-color: #3B82F6 !important; }
.main button:active { transform: translateY(1px) !important; }
.main [data-testid="stVerticalBlockBorderWrapper"] { border-radius: 14px !important; }

@media (prefers-reduced-motion: reduce) {
    *, *::before, *::after { scroll-behavior: auto !important; transition-duration: .01ms !important; animation-duration: .01ms !important; }
}

@media (max-width: 700px) {
    .app-hero, .auth-hero { padding: 1.25rem; }
    .brand-mark { width:36px; height:36px; margin-right:.75rem; }
}

    </style>
    """,
    unsafe_allow_html=True
)


# ============================================================
# SESSION STATE
# ============================================================

defaults = {

    # ========================================================
    # AUTHENTICATION
    # ========================================================

    "authenticated": False,

    "user": None,

    "chats_loaded": False,

    "auth_screen": "login",

    "password_reset_mode": False,

    "password_recovery_verified": False,

    "show_forgot_password": False,


    # ========================================================
    # STUDYSCAN STATE
    # ========================================================

    "generated_output": None,

    "chat_history": [],

    "active_chat_id": None,

    "current_feature": None,

    "current_study_level": None,

    "current_language": None,

    "current_file_names": [],


    # ========================================================
    # ASK STUDYSCAN
    # ========================================================

    "ask_studyscan_open": False,

    "ask_history": [],


    # ========================================================
    # CHANGE / SAVE CONTROL
    # ========================================================

    "pending_action": None,

    "pending_value": None,

    "pending_file_names": None,

    "current_chat_saved": False,


    # ========================================================
    # UPLOADER
    # ========================================================

    "uploader_version": 0,

    "ask_question_value": "",


    # ========================================================
    # MODE
    # ========================================================

    "current_mode": "Study Material",

    "document_type": None,

    "document_level": None,

    "theme": "System default",

}


for key, value in defaults.items():

    if key not in st.session_state:

        st.session_state[key] = value

    if st.session_state.get("theme") != st.session_state.get(
        "theme_selector",
        st.session_state.get("theme")
    ):

        st.session_state.theme = st.session_state.get(
            "theme_selector",
            st.session_state.theme
        )

# ============================================================
# PASSWORD RECOVERY
# ============================================================

recovery_token = st.query_params.get(
    "token_hash"
)

recovery_type = st.query_params.get(
    "type"
)


if (
    recovery_token
    and
    recovery_type == "recovery"
    and
    not st.session_state.password_recovery_verified
):

    (
        recovery_success,
        recovery_user,
        recovery_message
    ) = verify_password_recovery(
        recovery_token
    )


    # Remove the sensitive token from the browser URL
    st.query_params.clear()


    if recovery_success:

        st.session_state.password_reset_mode = True

        st.session_state.password_recovery_verified = True

        st.session_state.authenticated = True

        st.session_state.user = recovery_user

        st.session_state.chats_loaded = True

    else:

        st.session_state.password_reset_mode = False

        st.session_state.password_recovery_verified = False

        st.error(
            "This password reset link is invalid or has expired."
        )

        st.stop()

# ============================================================
# RESTORE AUTHENTICATION
# ============================================================

if not st.session_state.authenticated:

    current_user = get_current_user()

    if current_user is not None:

        st.session_state.authenticated = True

        st.session_state.user = current_user




# ============================================================
# PASSWORD RESET SCREEN
# ============================================================

if st.session_state.password_reset_mode:

    st.markdown(
        """
        <div style="
            text-align:center;
            padding-top:80px;
            padding-bottom:20px;
        ">
            <h1>🔐 Reset Your Password</h1>
            <p style="font-size:18px;">
                Create a new password for your StudyScan AI account.
            </p>
        </div>
        """,
        unsafe_allow_html=True
    )


    new_password = st.text_input(
        "New Password",
        type="password",
        placeholder="Enter your new password",
        key="reset_new_password"
    )


    confirm_new_password = st.text_input(
        "Confirm New Password",
        type="password",
        placeholder="Re-enter your new password",
        key="reset_confirm_password"
    )


    if st.button(
        "Update Password",
        use_container_width=True,
        key="update_password_button"
    ):

        if not new_password or not confirm_new_password:

            st.warning(
                "Please fill in both password fields."
            )

        elif new_password != confirm_new_password:

            st.error(
                "Passwords do not match."
            )

        elif len(new_password) < 6:

            st.error(
                "Password must contain at least 6 characters."
            )

        else:

            with st.spinner(
                "Updating your password..."
            ):

                success, message = update_password(
                    new_password
                )


            if success:

                st.success(
                    "Password updated successfully."
                )

                st.session_state.password_reset_mode = False

                st.session_state.password_recovery_verified = False

                st.session_state.authenticated = False

                st.session_state.user = None

                st.session_state.chats_loaded = False

                st.rerun()

            else:

                st.error(message)


    st.stop()


# ============================================================
# AUTHENTICATION SCREEN
# ============================================================

if not st.session_state.authenticated:

    st.markdown(
        """
        <div class="auth-page-marker"></div>
        <div class="auth-hero">
            <div class="brand-mark">S</div>
            <div>
                <p class="eyebrow">STUDYSCAN AI</p>
                <h1>Study smarter, not harder.</h1>
                <p>Turn notes and documents into focused study material in a few clicks.</p>
            </div>
        </div>
        """,
        unsafe_allow_html=True
    )


    # ========================================================
    # LOGIN / SIGN UP TABS
    # ========================================================

    login_tab, signup_tab = st.tabs(
        [
            ":material/login: Sign in",
            ":material/person_add: Create account"
        ]
    )


    # ========================================================
    # LOGIN
    # ========================================================

    with login_tab:

        st.subheader("Welcome back")
        st.caption("Sign in to keep your study materials together.")

        login_email = st.text_input(
            "Email",
            placeholder="Enter your email",
            key="login_email"
        )

        login_password = st.text_input(
            "Password",
            type="password",
            placeholder="Enter your password",
            key="login_password"
        )

        if st.button(
            "Forgot Password?",
            key="forgot_password_button"
        ):

            st.session_state.show_forgot_password = True


        if st.session_state.get(
            "show_forgot_password",
            False
        ):

            st.info(
                "Enter your account email and we will "
                "send you a password reset link."
            )

            forgot_email = st.text_input(
                "Account Email",
                placeholder="Enter your registered email",
                key="forgot_email"
            )

            if st.button(
                "Send Reset Link",
                use_container_width=True,
                key="send_reset_link"
            ):

                success, message = forgot_password(
                    forgot_email
                )

                if success:

                    st.success(message)

                else:

                    st.error(message)

            if st.button(
                "Cancel",
                key="cancel_forgot_password"
            ):

                st.session_state.show_forgot_password = False

                st.rerun()


        if st.button(
            "Sign in",
            width="stretch",
            type="primary",
            key="login_button"
        ):

            if not login_email or not login_password:

                st.warning(
                    "Please enter both email and password."
                )

            else:

                with st.spinner("Logging you in..."):

                    success, message = login(
                        login_email,
                        login_password
                    )


                if success:

                    st.session_state.authenticated = True

                    st.session_state.chats_loaded = False

                    st.success(
                        "Login successful."
                    )

                    st.rerun()

                else:

                    st.error(message)


    # ========================================================
    # SIGN UP
    # ========================================================

    with signup_tab:

        st.subheader("Create your account")
        st.caption("Save your generated materials and pick up where you left off.")

        signup_email = st.text_input(
            "Email",
            placeholder="Enter your email",
            key="signup_email"
        )

        signup_password = st.text_input(
            "Password",
            type="password",
            placeholder="Create a password",
            key="signup_password"
        )

        signup_confirm_password = st.text_input(
            "Confirm Password",
            type="password",
            placeholder="Re-enter your password",
            key="signup_confirm_password"
        )


        if st.button(
            "Create account",
            width="stretch",
            type="primary",
            key="signup_button"
        ):

            if (
                not signup_email
                or
                not signup_password
                or
                not signup_confirm_password
            ):

                st.warning(
                    "Please fill in all fields."
                )

            elif signup_password != signup_confirm_password:

                st.error(
                    "Passwords do not match."
                )

            else:

                with st.spinner(
                    "Creating your account..."
                ):

                    success, message = sign_up(
                        signup_email,
                        signup_password
                    )


                if success:

                    st.success(message)

                    st.info(
                        "After verifying your email, "
                        "return to the Login tab and log in."
                    )

                else:

                    st.error(message)


    # ========================================================
    # STOP THE REST OF THE APPLICATION
    # ========================================================

    st.stop()

# ============================================================
# LOAD USER CHATS FROM SUPABASE
# ============================================================

if (
    st.session_state.authenticated
    and
    st.session_state.user is not None
    and
    not st.session_state.get(
        "chats_loaded",
        False
    )
):

    try:

        user_id = (
            st.session_state.user.id
        )

        database_chats = (
            load_user_chats(user_id)
        )


        converted_chats = []


        for chat in database_chats:

            converted_chats.append(
                {
                    "id":
                        chat["id"],

                    "title":
                        chat["title"],

                    "feature":
                        chat["feature"],

                    "study_level":
                        chat["study_level"],

                    "language":
                        chat["language"],

                    "file_names":
                        chat["file_names"]
                        or [],

                    "content":
                        chat["content"],

                    "timestamp":
                        chat.get(
                            "created_at",
                            ""
                        )
                }
            )


        st.session_state.chat_history = (
            converted_chats
        )

        st.session_state.chats_loaded = True


    except Exception as e:

        st.session_state.chats_loaded = True

        print(
            "Could not load saved chats:",
            e
        )

# ============================================================
# CLEAN OUTPUT
# ============================================================

def clean_output(text):

    if not text:
        return ""

    text = text.replace("\r\n", "\n")

    text = re.sub(
        r"^\s*#{1,6}\s*",
        "",
        text,
        flags=re.MULTILINE
    )

    text = text.replace("**", "")
    text = text.replace("__", "")
    text = text.replace("*", "")
    text = text.replace("`", "")

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    return text.strip()


# ============================================================
# UNSAVED MATERIAL CHECK
# ============================================================

def has_unsaved_material():

    return (
        st.session_state.generated_output
        is not None
        and
        st.session_state.generated_output.strip()
        != ""
        and
        not st.session_state.current_chat_saved
    )


# ============================================================
# FILE NAME
# ============================================================

def get_display_file_name():

    files = st.session_state.current_file_names

    if not files:

        return "Study Material"

    if len(files) == 1:

        return files[0]

    return (
        f"{files[0]} + "
        f"{len(files) - 1} more"
    )


# ============================================================
# SAVED CHAT TITLE
# ============================================================

def create_chat_title():

    feature = (
        st.session_state.current_feature
        if st.session_state.current_feature
        else "Not selected"
    )

    level = (
        st.session_state.current_study_level
        if st.session_state.current_study_level
        else "Not selected"
    )

    language = (
        st.session_state.current_language
        if st.session_state.current_language
        else "Not selected"
    )

    return (
        f"{get_display_file_name()} - "
        f"{feature} • "
        f"{level} • "
        f"{language}"
    )


# ============================================================
# SAVE CURRENT MATERIAL
# ============================================================

def save_current_material():

    if not has_unsaved_material():

        return False


    user = st.session_state.get(
        "user"
    )


    if not user:

        return False


    user_id = user.id


    title = create_chat_title()


    content = clean_output(
        st.session_state.generated_output
    )


    try:

        saved_chat = save_chat_to_database(

            user_id=user_id,

            title=title,

            feature=(
                st.session_state.current_feature
            ),

            study_level=(
                st.session_state.current_study_level
            ),

            language=(
                st.session_state.current_language
            ),

            file_names=list(
                st.session_state.current_file_names
            ),

            content=content
        )


        if not saved_chat:

            return False


        # ----------------------------------------------------
        # KEEP LOCAL COPY FOR IMMEDIATE SIDEBAR DISPLAY
        # ----------------------------------------------------

        chat = {

            "id":
                saved_chat["id"],

            "title":
                title,

            "feature":
                st.session_state.current_feature,

            "study_level":
                st.session_state.current_study_level,

            "language":
                st.session_state.current_language,

            "file_names":
                list(
                    st.session_state.current_file_names
                ),

            "content":
                content,

            "timestamp":
                saved_chat.get(
                    "created_at",
                    datetime.now().strftime(
                        "%d %b %Y, %I:%M %p"
                    )
                )
        }


        st.session_state.chat_history.insert(
            0,
            chat
        )


        st.session_state.active_chat_id = (
            saved_chat["id"]
        )


        st.session_state.current_chat_saved = True


        return True


    except Exception as e:

        print(
            "Could not save chat:",
            e
        )

        st.session_state.current_chat_saved = False

        return False


# ============================================================
# OPEN SAVED CHAT
# ============================================================

def open_saved_chat(chat):

    st.session_state.generated_output = (
        chat["content"]
    )

    st.session_state.current_feature = (
        chat["feature"]
    )

    st.session_state.current_study_level = (
        chat["study_level"]
    )

    st.session_state.current_language = (
        chat["language"]
    )

    st.session_state.current_file_names = (
        list(chat["file_names"])
    )

    st.session_state.active_chat_id = (
        chat["id"]
    )

    st.session_state.ask_history = []

    st.session_state.ask_studyscan_open = False

    st.session_state.ask_question = ""

    if chat["feature"] == "Read & Understand Document":

        st.session_state.current_mode = "Documents"

        st.session_state.document_level = (
            chat["study_level"]
        )

    else:

        st.session_state.current_mode = "Study Material"


# ============================================================
# APPLY PENDING ACTION
# ============================================================

def apply_pending_action():

    action = st.session_state.pending_action

    value = st.session_state.pending_value


    if action == "mode":

        st.session_state.current_mode = value

        st.session_state.generated_output = None

        st.session_state.current_feature = None

        st.session_state.current_study_level = None

        st.session_state.current_language = None

        st.session_state.current_file_names = []

        st.session_state.ask_history = []

        st.session_state.ask_studyscan_open = False

        st.session_state.uploader_version += 1


    elif action == "feature":

        st.session_state.current_feature = value

        st.session_state.generated_output = None

        st.session_state.ask_history = []


    elif action == "level":

        st.session_state.current_study_level = value

        st.session_state.generated_output = None

        st.session_state.ask_history = []


    elif action == "language":

        st.session_state.current_language = value

        st.session_state.generated_output = None

        st.session_state.ask_history = []


    elif action == "files":

        st.session_state.current_file_names = (
            list(
                st.session_state.pending_file_names
            )
        )

        st.session_state.generated_output = None

        st.session_state.ask_history = []


    elif action == "open_chat":

        chat = next(
            (
                item
                for item in st.session_state.chat_history
                if item["id"] == value
            ),
            None
        )

        if chat:

            open_saved_chat(chat)


    elif action == "new_chat":

        st.session_state.generated_output = None

        st.session_state.current_chat_saved = False

        st.session_state.chats_loaded = False

        st.session_state.active_chat_id = None

        st.session_state.current_feature = None

        st.session_state.current_study_level = None

        st.session_state.current_language = None

        st.session_state.current_file_names = []

        st.session_state.ask_history = []

        st.session_state.ask_studyscan_open = False

        st.session_state.ask_question = ""

        st.session_state.current_mode = "Study Material"

        st.session_state.document_type = "General Document"

        st.session_state.document_level = "Moderate"

        st.session_state.uploader_version += 1


    st.session_state.pending_action = None

    st.session_state.pending_value = None

    st.session_state.pending_file_names = None


# ============================================================
# SAVE CONFIRMATION DIALOG
# ============================================================

@st.dialog("Unsaved Material")
def save_confirmation_dialog():

    st.markdown(
        """
        You have generated material that has not been saved.

        Would you like to save it before continuing?
        """
    )

    st.write("")

    col1, col2 = st.columns(2)


    with col1:

        if st.button(
            "💾 Save & Continue",
            use_container_width=True
        ):

            save_current_material()

            apply_pending_action()

            st.rerun()


    with col2:

        if st.button(
            "🗑️ Don't Save",
            use_container_width=True
        ):

            apply_pending_action()

            st.rerun()


    st.write("")


    if st.button(
        "↩️ Cancel",
        use_container_width=True
    ):

        st.session_state.pending_action = None

        st.session_state.pending_value = None

        st.session_state.pending_file_names = None

        st.rerun()


# ============================================================
# REQUEST CHANGE
# ============================================================

def request_change(
    action,
    value=None,
    file_names=None
):

    st.session_state.pending_action = action

    st.session_state.pending_value = value

    st.session_state.pending_file_names = file_names


    if has_unsaved_material():

        save_confirmation_dialog()

    else:

        apply_pending_action()

        st.rerun()


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.markdown("<div class='sidebar-brand'><span> S </span><strong>StudyScan</strong></div>", unsafe_allow_html=True)
    st.caption("Your learning workspace")


    # ========================================================
    # NEW CHAT
    # ========================================================

    if st.button(
        "New study session",
        icon=":material/add:",
        width="stretch",
        type="primary",
    ):

        request_change(
            "new_chat"
        )


    st.markdown("### Recent sessions")


    # ========================================================
    # SAVED CHATS
    # ========================================================

    if not st.session_state.chat_history:

        st.caption(
            "No saved study material yet."
        )

    else:

        for chat in st.session_state.chat_history:

            chat_col, delete_col = st.columns(
                [5, 1]
            )

            with chat_col:

                if st.button(
                    f"📄 {chat['title']}",
                    key=f"saved_chat_{chat['id']}",
                    use_container_width=True
                ):

                    request_change(
                        "open_chat",
                        chat["id"]
                    )

            with delete_col:

                if st.button(
                    "🗑️",
                    key=f"delete_chat_{chat['id']}",
                    help="Delete this chat"
                ):

                    current_user = (
                        st.session_state.get("user")
                    )

                    if current_user:

                        success = (
                            delete_chat_from_database(
                                chat["id"],
                                current_user.id
                            )
                        )

                        if success:

                            st.session_state.chat_history = [
                                item
                                for item in st.session_state.chat_history
                                if item["id"] != chat["id"]
                            ]

                            if (
                                st.session_state.active_chat_id
                                == chat["id"]
                            ):

                                st.session_state.generated_output = None

                                st.session_state.active_chat_id = None
 
                                st.session_state.current_feature = None

                                st.session_state.current_study_level = None

                                st.session_state.current_language = None

                                st.session_state.current_file_names = []

                            st.rerun()

                        else:

                            st.error(
                                "Could not delete this chat."
                            )


    st.markdown("---")


    










    # ========================================================
    # ASK STUDYSCAN
    # ========================================================

    if st.button(
        "Ask about this material",
        icon=":material/auto_awesome:",
        width="stretch",
    ):

        if not st.session_state.generated_output:

            st.warning(
                "Generate material first."
            )

        else:

            st.session_state.ask_studyscan_open = True

            st.rerun()


# ============================================================
# MAIN PAGE
# ============================================================


# ============================================================
# STUDYSCAN TOP-RIGHT MENU
# ============================================================

top_menu = st.popover(
    "⋮",
    type="secondary",
    width=260,
    key="studyscan_top_menu"
)

with top_menu:

    st.markdown("### StudyScan AI")

    # ========================================================
    # THEME
    # ========================================================

    st.markdown("**🎨 Theme**")

    selected_theme = st.radio(
        "Choose theme",
        [
            "Dark",
            "Light"
        ],
        index=0,
        key="studyscan_top_menu_theme",
        label_visibility="collapsed"
    )

    st.markdown("---")

    # ========================================================
    # DELETE CHAT HISTORY
    # ========================================================

    if st.button(
        "🗑️ Delete Chat History",
        use_container_width=True,
        key="top_delete_history"
    ):

        current_user = (
            st.session_state.get("user")
        )

        if current_user:

            success = delete_all_user_chats(
                current_user.id
            )

            if success:

                st.session_state.chat_history = []

                st.session_state.generated_output = None

                st.session_state.active_chat_id = None

                st.session_state.current_feature = None

                st.session_state.current_study_level = None

                st.session_state.current_language = None

                st.session_state.current_file_names = []

                st.session_state.current_chat_saved = False

                st.session_state.ask_history = []

                st.session_state.ask_studyscan_open = False

                st.rerun()

            else:

                st.error(
                    "Could not delete chat history."
                )

    st.markdown("---")

    # ========================================================
    # SIGN OUT
    # ========================================================

    if st.button(
        "🚪 Sign Out",
        use_container_width=True,
        key="top_sign_out"
    ):

        logout()

        st.rerun()




st.markdown(
    """
    <div class="app-hero">
        <div>
            <p class="eyebrow">YOUR LEARNING WORKSPACE</p>
            <h1>Make every page easier to learn.</h1>
            <p>Upload notes, choose an outcome, and get study-ready material tailored to you.</p>
        </div>
        <div class="hero-orb">✦</div>
    </div>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# MODE SELECTION
# ============================================================

st.markdown("## Start with a workspace")
st.caption("Switch between study material and document understanding at any time.")

mode_options = ["Study material", "Documents"]


current_mode_display = (
    "Study material"
    if st.session_state.current_mode == "Study Material"
    else "Documents"
)


selected_mode_display = st.segmented_control(
    "Choose a workspace",
    mode_options,
    default=current_mode_display,
    required=True,
    width="stretch",
    key="workspace_mode",
)


selected_mode = (
    "Study Material"
    if selected_mode_display == "Study material"
    else "Documents"
)


if selected_mode != st.session_state.current_mode:

    request_change(
        "mode",
        selected_mode
    )


st.markdown("---")


# ################################################################
# ################################################################
# STUDY MATERIAL WORKFLOW
# ################################################################
# ################################################################

if st.session_state.current_mode == "Study Material":

    st.markdown(
        "## Create study material"
    )
    st.caption("Add your notes, choose the learning outcome, then generate a polished study guide.")


    # ============================================================
    # STUDY MATERIAL FILE UPLOAD
    # ============================================================

    uploaded_files = st.file_uploader(
        "Upload your study material",
        type=[
            "jpg",
            "jpeg",
            "png",
            "pdf"
        ],
        accept_multiple_files=True,
        key=f"study_uploader_{st.session_state.uploader_version}"
    )


    uploaded_names = sorted(
        [
            file.name
            for file in uploaded_files
        ]
        if uploaded_files
        else []
    )


    previous_names = sorted(
        st.session_state.current_file_names
    )


    if (
        uploaded_names
        and
        uploaded_names != previous_names
    ):

        if has_unsaved_material():

            request_change(
                "files",
                file_names=uploaded_names
            )

        else:

            st.session_state.current_file_names = (
                uploaded_names
            )


    # ============================================================
    # FEATURE
    # ============================================================

    feature_options = [

        "Summary",

        "Detailed Explanation",

        "Flashcards",

        "Exam Questions",

        "MCQs",

        "One-word Questions",

        "Exam Quiz"

    ]


    selected_feature = st.selectbox(
        "What would you like to generate?",

        feature_options,

        index=(
            feature_options.index(
                st.session_state.current_feature
            )
            if st.session_state.current_feature
            in feature_options
            else None
        ),

        placeholder="Select the type of study material"
    )


    if (
        selected_feature is not None
        and
        selected_feature
        != st.session_state.current_feature
    ):

        request_change(
            "feature",
            selected_feature
        )


    # ============================================================
    # STUDY LEVEL
    # ============================================================

    study_levels = [

        "School",

        "College",

        "Competitive Exam"

    ]


    selected_level = st.selectbox(
        "Study Level",

        study_levels,

        index=(
            study_levels.index(
                st.session_state.current_study_level
            )
            if st.session_state.current_study_level
            in study_levels
            else None
        ),

        placeholder="Select the level of study material"
    )


    if (
        selected_level is not None
        and
        selected_level
        != st.session_state.current_study_level
    ):

        request_change(
            "level",
            selected_level
        )


    # ============================================================
    # LANGUAGE
    # ============================================================

    languages = [

        "English",

        "Hindi",

        "Marathi",

        "Bengali",

        "Tamil",

        "Telugu",

        "Gujarati",

        "Kannada",

        "Malayalam",

        "Punjabi",

        "Urdu"

    ]


    selected_language = st.selectbox(
        "Output Language",

        languages,

        index=(
            languages.index(
                st.session_state.current_language
            )
            if st.session_state.current_language
            in languages
            else None
        ),

        placeholder="Select the language"
    )


    if (
        selected_language is not None
        and
        selected_language
        != st.session_state.current_language
    ):

        request_change(
            "language",
            selected_language
        )


    # ============================================================
    # GENERATE STUDY MATERIAL
    # ============================================================

    generate_disabled = not (
        uploaded_files
        and
        selected_feature is not None
        and
        selected_level is not None
        and
        selected_language is not None
    )


    if st.button(
        "Generate study material",
        icon=":material/auto_awesome:",
        width="stretch",
        type="primary",
        disabled=generate_disabled
    ):

        if not uploaded_files:

            st.warning(
                "Please upload at least one "
                "study-material file first."
            )

            st.stop()


        if selected_feature is None:

            st.warning(
                "Please select the type of study material."
            )

            st.stop()


        if selected_level is None:

            st.warning(
                "Please select the level of study material."
            )

            st.stop()


        if selected_language is None:

            st.warning(
                "Please select the output language."
            )

            st.stop()


        # ========================================================
        # CREATE PROMPT
        # ========================================================

        prompt = get_prompt(
            selected_feature,
            selected_level,
            selected_language
        )


        # ========================================================
        # COUNT PAGES
        # ========================================================

        total_pages = 0

        image_files = []

        pdf_files = []


        for uploaded_file in uploaded_files:

            filename = uploaded_file.name.lower()


            if filename.endswith(
                (".jpg", ".jpeg", ".png")
            ):

                image_files.append(
                    uploaded_file
                )

                total_pages += 1


            elif filename.endswith(".pdf"):

                pdf_files.append(
                    uploaded_file
                )

                try:

                    total_pages += (
                        get_pdf_page_count(
                            uploaded_file
                        )
                    )

                except Exception as e:

                    st.error(
                        f"Could not read "
                        f"{uploaded_file.name}"
                    )

                    st.code(
                        str(e)
                    )

                    st.stop()


        # ========================================================
        # MAX PAGE LIMIT
        # ========================================================

        MAX_PAGES = 100


        if total_pages > MAX_PAGES:

            st.error(
                f"This upload contains "
                f"{total_pages} pages. "
                f"The current prototype supports "
                f"a maximum of {MAX_PAGES} combined pages."
            )

            st.stop()


        st.info(
            f"📚 StudyScan received "
            f"{total_pages} page(s) from "
            f"{len(uploaded_files)} file(s)."
        )


        # ========================================================
        # PREPARE IMAGE BATCHES
        # ========================================================

        with st.spinner(
            "📚 Preparing your study material..."
        ):

            try:

                image_batches = []


                for uploaded_file in image_files:

                    image = Image.open(
                        uploaded_file
                    )


                    if image.mode != "RGB":

                        image = image.convert(
                            "RGB"
                        )


                    image_batches.append(
                        [image]
                    )


                for uploaded_pdf in pdf_files:

                    pdf_batches = (
                        pdf_to_image_batches(
                            uploaded_pdf
                        )
                    )

                    image_batches.extend(
                        pdf_batches
                    )


                if not image_batches:

                    st.error(
                        "No readable study pages were found."
                    )

                    st.stop()


            except Exception as e:

                st.error(
                    "Something went wrong while "
                    "preparing your files."
                )

                st.code(
                    str(e)
                )

                st.stop()


        # ========================================================
        # BATCH PROCESSING
        # ========================================================

        batch_results = []

        total_batches = len(
            image_batches
        )


        progress_bar = st.progress(0)

        status_text = st.empty()


        for batch_number, batch_images in enumerate(
            image_batches,
            start=1
        ):

            success = False

            max_attempts = 3


            for attempt in range(
                1,
                max_attempts + 1
            ):

                status_text.write(
                    f"🧠 Analysing batch "
                    f"{batch_number} of "
                    f"{total_batches} "
                    f"(attempt "
                    f"{attempt}/"
                    f"{max_attempts})..."
                )


                try:

                    batch_result = (
                        generate_response(
                            prompt,
                            batch_images
                        )
                    )


                    batch_results.append(
                        batch_result
                    )

                    success = True

                    break


                except Exception as e:

                    error_text = str(e)


                    if (
                        "503" in error_text
                        or
                        "UNAVAILABLE" in error_text
                    ):

                        if attempt < max_attempts:

                            status_text.write(
                                "⚠️ Gemini is temporarily "
                                "unavailable. Retrying..."
                            )

                            time.sleep(
                                5 * attempt
                            )

                        else:

                            st.error(
                                "Gemini remained unavailable "
                                "after multiple attempts."
                            )

                            st.code(
                                error_text
                            )


                    else:

                        st.error(
                            f"Batch {batch_number} "
                            f"could not be processed."
                        )

                        st.code(
                            error_text
                        )


            if not success:

                st.stop()


            progress_bar.progress(
                batch_number / total_batches
            )


        # ========================================================
        # COMBINE RESULTS
        # ========================================================

        status_text.write(
            "🧠 Combining the analysed study material..."
        )


        combined_result = "\n\n".join(
            batch_results
        )


        cleaned_result = clean_output(
            combined_result
        )


        # ========================================================
        # SAVE GENERATED RESULT
        # ========================================================

        st.session_state.generated_output = (
            cleaned_result
        )


        st.session_state.current_feature = (
            selected_feature
        )

        st.session_state.current_study_level = (
            selected_level
        )

        st.session_state.current_language = (
            selected_language
        )

        st.session_state.current_file_names = (
            uploaded_names
        )


        # ========================================================
        # SAVE GENERATED STUDY MATERIAL
        # ========================================================

        st.session_state.current_chat_saved = False


        if not save_current_material():

            st.warning(
                "Study material was generated, "
                "but could not be saved to your account."
            )
        st.session_state.ask_history = []

        st.session_state.ask_studyscan_open = False

        st.session_state.ask_question = ""
        
        status_text.empty()

        progress_bar.empty()

        st.rerun()


# ################################################################
# ################################################################
# DOCUMENT WORKFLOW
# ################################################################
# ################################################################

if st.session_state.current_mode == "Documents":

    st.markdown(
        "## 📄 Read & Understand Documents"
    )

    st.markdown(
        "Upload an official or formal document and "
        "StudyScan AI will explain its contents in "
        "clear and understandable language while "
        "preserving important information."
    )


    # ============================================================
    # DOCUMENT TYPE
    # ============================================================

    document_types = [

        "General Document",

        "School / College Notice",

        "Government Document",

        "Affidavit",

        "Legal Notice",

        "Court Summons",

        "Agreement / Contract",

        "Property Document",

        "Financial Document",

        "Other"

    ]


    selected_document_type = st.selectbox(
        "Document Type",
        ["Select document type"] + document_types,
        index=0,
    )


    if (
        selected_document_type
        != st.session_state.document_type
    ):

        st.session_state.document_type = (
            selected_document_type
        )

        st.session_state.generated_output = None

        st.session_state.ask_history = []


    # ============================================================
    # DOCUMENT LEVEL
    # ============================================================

    document_levels = [

        "Easy",

        "Moderate",

        "Tough"

    ]


    selected_document_level = st.selectbox(
        "Explanation Level",
         document_levels,
        index= None,
        placeholder="Select explanation level",

         key="document_level_selector"
    )


    if (
        selected_document_level != "Select explanation level"
        and
        selected_document_level
        != st.session_state.document_level
    ):

        st.session_state.document_level = (
            selected_document_level
        )

        st.session_state.generated_output = None

        st.session_state.ask_history = []


    # ============================================================
    # DOCUMENT LANGUAGE
    # ============================================================

    document_languages = [

        "English",

        "Hindi",

        "Marathi",

        "Bengali",

        "Tamil",

        "Telugu",

        "Gujarati",

        "Kannada",

        "Malayalam",

        "Punjabi",

        "Urdu"

    ]


    selected_document_language = st.selectbox(
        "Output Language",

        document_languages,

        index=(
            document_languages.index(
                st.session_state.current_language
            )
            if st.session_state.current_language
            in document_languages
            else None
        ),

        placeholder="Select the language",

        key="document_language_selector"
    )


    if (
        selected_document_language is not None
        and
        selected_document_language
        != st.session_state.current_language
    ):

        st.session_state.current_language = (
            selected_document_language
        )

        st.session_state.generated_output = None

        st.session_state.ask_history = []


    # ============================================================
    # DOCUMENT UPLOAD
    # ============================================================

    st.markdown(
        "### Upload Document"
    )


    document_files = st.file_uploader(
        "Upload your document",
        type=[
            "jpg",
            "jpeg",
            "png",
            "pdf"
        ],
        accept_multiple_files=True,
        key=f"document_uploader_{st.session_state.uploader_version}"
    )


    document_names = sorted(
        [
            file.name
            for file in document_files
        ]
        if document_files
        else []
    )


    if (
        document_names
        and
        document_names
        != sorted(
            st.session_state.current_file_names
        )
    ):

        if has_unsaved_material():

            request_change(
                "files",
                file_names=document_names
            )

        else:

            st.session_state.current_file_names = (
                document_names
            )


    # ============================================================
    # GENERATE DOCUMENT SUMMARY
    # ============================================================

    document_generate_disabled = not (
        document_files
        and
        selected_document_language is not None
    )


    if st.button(
        "📄 Generate Document Summary",
        use_container_width=True,
        disabled=document_generate_disabled,
        key="generate_document"
    ):

        if not document_files:

            st.warning(
                "Please upload at least one document first."
            )

            st.stop()


        if selected_document_language is None:

            st.warning(
                "Please select the output language."
            )

            st.stop()


        # ========================================================
        # CREATE DOCUMENT PROMPT
        # ========================================================

        document_prompt = get_prompt(
            "Read & Understand Document",
            selected_document_level,
            selected_document_language,
            selected_document_type
        )


        # ========================================================
        # COUNT DOCUMENT PAGES
        # ========================================================

        total_pages = 0

        document_image_files = []

        document_pdf_files = []


        for uploaded_file in document_files:

            filename = uploaded_file.name.lower()


            if filename.endswith(
                (".jpg", ".jpeg", ".png")
            ):

                document_image_files.append(
                    uploaded_file
                )

                total_pages += 1


            elif filename.endswith(".pdf"):

                document_pdf_files.append(
                    uploaded_file
                )

                try:

                    total_pages += (
                        get_pdf_page_count(
                            uploaded_file
                        )
                    )

                except Exception as e:

                    st.error(
                        f"Could not read "
                        f"{uploaded_file.name}"
                    )

                    st.code(
                        str(e)
                    )

                    st.stop()


        # ========================================================
        # MAX PAGE LIMIT
        # ========================================================

        MAX_PAGES = 100


        if total_pages > MAX_PAGES:

            st.error(
                f"This document contains "
                f"{total_pages} pages. "
                f"The current prototype supports "
                f"a maximum of {MAX_PAGES} combined pages."
            )

            st.stop()


        st.info(
            f"📄 StudyScan received "
            f"{total_pages} page(s) from "
            f"{len(document_files)} file(s)."
        )


        # ========================================================
        # PREPARE DOCUMENT IMAGE BATCHES
        # ========================================================

        with st.spinner(
            "📄 Reading and preparing your document..."
        ):

            try:

                document_batches = []


                # ------------------------------------------------
                # IMAGE FILES
                # ------------------------------------------------

                for uploaded_file in document_image_files:

                    image = Image.open(
                        uploaded_file
                    )


                    if image.mode != "RGB":

                        image = image.convert(
                            "RGB"
                        )


                    document_batches.append(
                        [image]
                    )


                # ------------------------------------------------
                # PDF FILES
                # ------------------------------------------------

                for uploaded_pdf in document_pdf_files:

                    pdf_batches = (
                        pdf_to_image_batches(
                            uploaded_pdf
                        )
                    )

                    document_batches.extend(
                        pdf_batches
                    )


                if not document_batches:

                    st.error(
                        "No readable document pages were found."
                    )

                    st.stop()


            except Exception as e:

                st.error(
                    "Something went wrong while "
                    "preparing your document."
                )

                st.code(
                    str(e)
                )

                st.stop()


        # ========================================================
        # DOCUMENT BATCH PROCESSING
        # ========================================================

        document_results = []

        total_document_batches = len(
            document_batches
        )


        document_progress = st.progress(0)

        document_status = st.empty()


        for batch_number, batch_images in enumerate(
            document_batches,
            start=1
        ):

            success = False

            max_attempts = 3


            for attempt in range(
                1,
                max_attempts + 1
            ):

                document_status.write(
                    f"📄 Reading document batch "
                    f"{batch_number} of "
                    f"{total_document_batches} "
                    f"(attempt "
                    f"{attempt}/"
                    f"{max_attempts})..."
                )


                try:

                    batch_result = (
                        generate_response(
                            document_prompt,
                            batch_images
                        )
                    )


                    document_results.append(
                        batch_result
                    )

                    success = True

                    break


                except Exception as e:

                    error_text = str(e)


                    if (
                        "503" in error_text
                        or
                        "UNAVAILABLE" in error_text
                    ):

                        if attempt < max_attempts:

                            document_status.write(
                                "⚠️ Gemini is temporarily "
                                "unavailable. Retrying..."
                            )

                            time.sleep(
                                5 * attempt
                            )

                        else:

                            st.error(
                                "Gemini remained unavailable "
                                "after multiple attempts."
                            )

                            st.code(
                                error_text
                            )


                    else:

                        st.error(
                            f"Document batch {batch_number} "
                            f"could not be processed."
                        )

                        st.code(
                            error_text
                        )


            if not success:

                st.stop()


            document_progress.progress(
                batch_number /
                total_document_batches
            )


        # ========================================================
        # COMBINE DOCUMENT RESULTS
        # ========================================================

        document_status.write(
            "📄 Combining the analysed document..."
        )


        combined_document_result = "\n\n".join(
            document_results
        )


        cleaned_document_result = clean_output(
            combined_document_result
        )


        # ========================================================
        # SAVE DOCUMENT RESULT
        # ========================================================

        st.session_state.generated_output = (
            cleaned_document_result
        )


        st.session_state.current_feature = (
            "Read & Understand Document"
        )


        st.session_state.current_study_level = (
            selected_document_level
        )


        st.session_state.current_language = (
            selected_document_language
        )


        st.session_state.current_file_names = (
            document_names
        )


        st.session_state.document_type = (
            selected_document_type
        )


        st.session_state.document_level = (
            selected_document_level
        )

        # ========================================================
        # SAVE GENERATED DOCUMENT
        # ========================================================

        st.session_state.current_chat_saved = False


        if not save_current_material():

            st.warning(
                "Document was understood successfully, "
                "but could not be saved to your account."
            )


        st.session_state.ask_history = []

        st.session_state.ask_studyscan_open = False

        st.session_state.ask_question = ""


        document_status.empty()

        document_progress.empty()


        st.rerun()


# ################################################################
# ################################################################
# GENERATED OUTPUT
# ################################################################
# ################################################################

if st.session_state.generated_output:

    st.markdown("---")


    if st.session_state.current_mode == "Documents":

        st.success(
            "📄 Your document has been understood successfully!"
        )

        st.markdown(
            "## 📄 Document Understanding"
        )

        st.caption(
            f"Document Type: "
            f"{st.session_state.document_type}  •  "
            f"Level: "
            f"{st.session_state.document_level}  •  "
            f"Language: "
            f"{st.session_state.current_language}"
        )

    else:

        st.success(
            "✨ Your study material is ready!"
        )

        st.markdown(
            f"## 📄 "
            f"{st.session_state.current_feature}"
        )


    # ============================================================
    # MAIN OUTPUT
    # ============================================================

    st.markdown(
        st.session_state.generated_output
    )


    # ============================================================
    # EXPORT
    # ============================================================

    st.markdown("---")

    if st.session_state.current_mode == "Documents":

        st.markdown(
            "### 📤 Export Document Summary"
        )

    else:

        st.markdown(
            "### 📤 Export Study Material"
        )


    output_text = clean_output(
        st.session_state.generated_output
    )


    col1, col2, col3 = st.columns(3)


    # ============================================================
    # TXT
    # ============================================================

    with col1:

        st.download_button(
            "📝 Notepad (.txt)",

            data=output_text,

            file_name=(
                "StudyScan_Document_Summary.txt"
                if st.session_state.current_mode == "Documents"
                else "StudyScan_Study_Material.txt"
            ),

            mime="text/plain",

            use_container_width=True
        )


    # ============================================================
    # WORD
    # ============================================================

    with col2:

        try:

            from docx import Document


            document = Document()


            document.add_heading(
                "StudyScan AI",
                level=1
            )


            document.add_heading(
                (
                    "Document Understanding"
                    if st.session_state.current_mode == "Documents"
                    else st.session_state.current_feature
                ),
                level=2
            )


            for line in output_text.splitlines():

                if line.strip():

                    document.add_paragraph(
                        line.strip()
                    )


            word_buffer = BytesIO()


            document.save(
                word_buffer
            )


            word_buffer.seek(0)


            st.download_button(
                "📘 Word (.docx)",

                data=word_buffer,

                file_name=(
                    "StudyScan_Document_Summary.docx"
                    if st.session_state.current_mode == "Documents"
                    else "StudyScan_Study_Material.docx"
                ),

                mime=(
                    "application/vnd.openxmlformats-"
                    "officedocument.wordprocessingml.document"
                ),

                use_container_width=True
            )


        except ImportError:

            st.warning(
                "Word export requires python-docx."
            )


    # ============================================================
    # PDF
    # ============================================================

    with col3:

        try:

            from reportlab.lib.pagesizes import A4

            from reportlab.platypus import (
                SimpleDocTemplate,
                Paragraph,
                Spacer
            )

            from reportlab.lib.styles import (
                getSampleStyleSheet
            )


            pdf_buffer = BytesIO()


            pdf_document = SimpleDocTemplate(
                pdf_buffer,
                pagesize=A4
            )


            styles = getSampleStyleSheet()


            story = []


            story.append(
                Paragraph(
                    "StudyScan AI",
                    styles["Title"]
                )
            )


            story.append(
                Spacer(1, 12)
            )


            story.append(
                Paragraph(
                    (
                        "Document Understanding"
                        if st.session_state.current_mode == "Documents"
                        else st.session_state.current_feature
                    ),
                    styles["Heading2"]
                )
            )


            story.append(
                Spacer(1, 12)
            )


            for line in output_text.splitlines():

                if line.strip():

                    safe_line = (
                        line
                        .replace("&", "&amp;")
                        .replace("<", "&lt;")
                        .replace(">", "&gt;")
                    )


                    story.append(
                        Paragraph(
                            safe_line,
                            styles["BodyText"]
                        )
                    )


                    story.append(
                        Spacer(1, 6)
                    )


            pdf_document.build(
                story
            )


            pdf_buffer.seek(0)


            st.download_button(
                "📄 PDF",

                data=pdf_buffer,

                file_name=(
                    "StudyScan_Document_Summary.pdf"
                    if st.session_state.current_mode == "Documents"
                    else "StudyScan_Study_Material.pdf"
                ),

                mime="application/pdf",

                use_container_width=True
            )


        except ImportError:

            st.warning(
                "PDF export requires reportlab."
            )


# ============================================================
# ASK STUDYSCAN — FIXED RIGHT PANEL
# ============================================================

if (
    st.session_state.ask_studyscan_open
    and
    st.session_state.generated_output
):

    with st.container(
        key="ask_panel"
    ):

        # ====================================================
        # HEADER
        # ====================================================

        st.markdown(
            '<div class="ask-header-title">'
            '🤖 Ask StudyScan'
            '</div>',
            unsafe_allow_html=True
        )

        st.markdown(
            '<div class="ask-header-subtitle">'
            'Your AI study companion'
            '</div>',
            unsafe_allow_html=True
        )

        st.markdown("---")


        # ====================================================
        # CHAT MESSAGES
        # ====================================================

        if not st.session_state.ask_history:

            st.markdown(
                """
                <div class="ask-bot-row">
                    <div class="ask-bot-bubble">
                        <div class="ask-bot-label">
                            🤖 StudyScan
                        </div>
                        Hi! Ask me anything about your generated material.
                    </div>
                </div>
                """,
                unsafe_allow_html=True
            )

        else:

            for message in (
                st.session_state.ask_history
            ):

                question = (
                    message.get(
                        "question",
                        ""
                    )
                )

                answer = (
                    message.get(
                        "answer",
                        ""
                    )
                )


                # --------------------------------------------
                # USER MESSAGE
                # --------------------------------------------

                safe_question = (
                    question
                    .replace(
                        "&",
                        "&amp;"
                    )
                    .replace(
                        "<",
                        "&lt;"
                    )
                    .replace(
                        ">",
                        "&gt;"
                    )
                    .replace(
                        "\n",
                        "<br>"
                    )
                )


                st.markdown(
                    f"""
                    <div class="ask-user-row">
                        <div class="ask-user-bubble">
                            <div class="ask-user-label">
                                You
                            </div>
                            {safe_question}
                        </div>
                    </div>
                    """,
                    unsafe_allow_html=True
                )


                # --------------------------------------------
                # BOT MESSAGE
                # --------------------------------------------

                safe_answer = (
                    answer
                    .replace(
                        "&",
                        "&amp;"
                    )
                    .replace(
                        "<",
                        "&lt;"
                    )
                    .replace(
                        ">",
                        "&gt;"
                    )
                    .replace(
                        "\n",
                        "<br>"
                    )
                )


                st.markdown(
                    f"""
                    <div class="ask-bot-row">
                        <div class="ask-bot-bubble">
                            <div class="ask-bot-label">
                                🤖 StudyScan
                            </div>
                            {safe_answer}
                        </div>
                    </div>
                    """,
                    unsafe_allow_html=True
                )


        st.markdown("---")


        # ====================================================
        # QUESTION INPUT
        # ====================================================

        with st.form(
            key="ask_question_form",
            clear_on_submit=True
        ):

            question = st.text_area(

                "Ask your question",

                placeholder=
                    "Type your question here...",

                height=140,

                key="ask_question",

                label_visibility="collapsed"
            )


            st.markdown(
                '<div style="height:6px;"></div>',
                unsafe_allow_html=True
            )


            send_clicked = st.form_submit_button(
                "➤ Send",
                use_container_width=True
            )


        # ====================================================
        # SEND QUESTION
        # ====================================================

        if send_clicked:

            cleaned_question = (
                question.strip()
            )


            if not cleaned_question:

                st.warning(
                    "Please enter a question first."
                )


            else:

                with st.spinner(
                    "StudyScan is thinking..."
                ):

                    try:

                        # ------------------------------------
                        # BUILD AI CHAT HISTORY
                        # ------------------------------------

                        ai_history = []


                        for message in (
                            st.session_state.ask_history
                        ):

                            ai_history.append(
                                {
                                    "role":
                                        "user",

                                    "text":
                                        message[
                                            "question"
                                        ]
                                }
                            )


                            ai_history.append(
                                {
                                    "role":
                                        "assistant",

                                    "text":
                                        message[
                                            "answer"
                                        ]
                                }
                            )


                        # ------------------------------------
                        # ASK GEMINI
                        # ------------------------------------

                        answer = ask_studyscan(

                            cleaned_question,

                            st.session_state.generated_output,

                            ai_history
                        )


                        answer = clean_output(
                            answer
                        )


                        # ------------------------------------
                        # SAVE CONVERSATION
                        # ------------------------------------

                        st.session_state.ask_history.append(
                            {
                                "question":
                                    cleaned_question,

                                "answer":
                                    answer
                            }
                        )


                        st.rerun()


                    except Exception as e:

                        st.error(
                            "Something went wrong "
                            "while answering."
                        )

                        st.code(
                            str(e)
                        )


        # ====================================================
        # CHAT CONTROLS
        # ====================================================

        if st.session_state.ask_history:

            st.markdown("")


            if st.button(
                "🗑️ Clear Chat",

                key="ask_clear",

                use_container_width=True
            ):

                st.session_state.ask_history = []

                st.rerun()


        # ====================================================
        # CLOSE ASK STUDYSCAN
        # ====================================================

        if st.button(
            "✕ Close Ask StudyScan",

            key="ask_close",

            use_container_width=True
        ):

            st.session_state.ask_studyscan_open = False

            st.rerun()
